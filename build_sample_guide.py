#!/usr/bin/env python3
"""Three-page guide: which of Nicolas Picard's eight files can be shared as samples."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

GREEN = RGBColor(0x1B, 0x3A, 0x2F)
GREEN_MID = RGBColor(0x2D, 0x6A, 0x4F)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x3D, 0x4A, 0x45)
RULE = "1B3A2F"


def set_run_font(run, size=10, bold=False, color=INK, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = "Calibri"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(attr), "Calibri")


def add_text(p, text, **kwargs):
    r = p.add_run(text)
    set_run_font(r, **kwargs)
    return r


def para(doc, text="", size=10, bold=False, color=INK, space_before=0, space_after=3,
         italic=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.02
    if text:
        add_text(p, text, size=size, bold=bold, color=color, italic=italic)
    return p


def rich(doc, runs, space_before=0, space_after=3):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.02
    for text, kwargs in runs:
        add_text(p, text, **kwargs)
    return p


def heading(doc, text):
    p = para(doc, text, size=12, bold=True, color=GREEN, space_before=6, space_after=1)
    pPr = p._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:keepNext"))
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), RULE)
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def bullet(doc, lead, rest, space_after=1):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.0
    pf.left_indent = Cm(0.35)
    pf.first_line_indent = Cm(-0.35)
    add_text(p, "•  ", size=10, color=GREEN_MID, bold=True)
    if lead:
        add_text(p, lead, size=10, bold=True, color=INK)
    add_text(p, rest, size=10, color=INK)
    return p


def footer(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    add_text(p, "Writing samples to share  ·  ", size=8, color=MUTED)
    for instr in ("PAGE", "NUMPAGES"):
        if instr == "NUMPAGES":
            add_text(p, " / ", size=8, color=MUTED)
        run = p.add_run()
        set_run_font(run, size=8, color=MUTED)
        fld = OxmlElement("w:fldChar")
        fld.set(qn("w:fldCharType"), "begin")
        run._r.append(fld)
        run2 = p.add_run()
        set_run_font(run2, size=8, color=MUTED)
        code = OxmlElement("w:instrText")
        code.set(qn("xml:space"), "preserve")
        code.text = f" {instr} "
        run2._r.append(code)
        run3 = p.add_run()
        set_run_font(run3, size=8, color=MUTED)
        fld2 = OxmlElement("w:fldChar")
        fld2.set(qn("w:fldCharType"), "end")
        run3._r.append(fld2)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(1.35)
    sec.right_margin = Cm(1.35)
    sec.top_margin = Cm(1.15)
    sec.bottom_margin = Cm(1.15)
    sec.header_distance = Cm(0.4)
    sec.footer_distance = Cm(0.4)
    footer(sec)

    para(doc, "Writing samples to share", size=16, bold=True, color=GREEN, space_after=0)
    para(
        doc,
        "Nicolas Picard  ·  which of the eight files can go on the site or to a client  ·  28 October 2026",
        size=9,
        color=MUTED,
        space_after=4,
    )
    para(
        doc,
        "None of these eight files is the public sample the site needs. The two pieces still to write stay first: a short briefing on what AI can and cannot do inside carbon-accounting software, and a short public paper on where supplier-carbon programs break. The files below are credentials and models. They stay off the site until the conditions in this note are met. None of them is a French sample, so a French proof still has to be written.",
        space_after=3,
    )
    para(
        doc,
        "Four working assumptions. AstroFinance has not given written permission to publish any excerpt. The AmCham file is a draft: it still carries a deadline line and old Asia Foundation email addresses. The CSR brief is in IG Advisors’ voice. Avon and WEConnect have not cleared the supplier-diversity plan. A written yes on one file changes the line for that file only.",
        size=10,
        italic=True,
        color=MUTED,
        space_after=2,
    )

    heading(doc, "What can go on the site")

    bullet(
        doc,
        "1. Two new public pieces, still to write. ",
        "Format: a regulatory briefing, and a practitioner white paper. Advertise with the public sentence: “I write practitioner white papers and case studies for supply-chain climate software teams.” These are the only samples that match that sentence. Use public facts. Footer on each: this explains the product; it is not emissions, legal, or investment advice. The first paid case study, with the customer’s written yes, is the case-study sample. None of the eight files fills that slot.",
        space_after=2,
    )
    bullet(
        doc,
        "2. AstroFinance white paper, sections 6.1 to 6.3, after written consent. ",
        "Use Business Model (platform, who it is for, core services, revenue model, token product suite), Technology Strategy, and Tokenisation Process. Format: a product white paper. Advertise it second, as a product explanation for a real-asset platform. It is a sample of long product writing. It is a different offer from the supply-chain climate work. Before anyone sees it: take out yields, market-size figures, the optional platform token, the roadmap, and any live deal; keep AstroFinance’s own disclaimer; send them the excerpt before it goes up. The byline on the file says Head of Product. On a writing sample, use the credit line you want a client to see, and confirm it with them. The three short pieces already scoped — how the asset hub works, what an asset needs before it can be offered, and security token versus utility token — are cleaner public samples than a cut from this 37-page paper. Leave off the executive summary, the market chapters, the blockchain primer in section 5, the compliance-jurisdiction pages, the roadmap, and the team page.",
        space_after=2,
    )
    bullet(
        doc,
        "3. The published AmCham article, as a link, if a published version exists. ",
        "“Non-Tariff Barriers: The Sticking Point for SMEs in ASEAN,” bylined with Véronique Salze-Lozac’h for The Asia Foundation, written for the ASEAN 2015 context. Format: a policy article, the ancestor of a regulatory briefing. Advertise it below the fold, dated about 2014: an early published briefing that explains a rule for the firms it affects. Link the published URL. Leave this draft PDF off the site. Leave the piece in its own decade. A carbon rewrite of it would sit outside the niche.",
        space_after=2,
    )
    bullet(
        doc,
        "4. The SimAxis revenue-token blueprint, after written AstroFinance consent, and heavily cut. ",
        "SimAxis is labeled fictional. The file is still AstroFinance confidential method. Share the section list — executive summary, company and deal overview, asset onboarding and ownership, communication and marketing planning, issuance, monitoring — plus a shortened executive summary. Take out the issuance partner, the fleet counts, the ROI scenarios, and any sentence that reads like an offer of a security. Format: a structured client memo. Advertise it as a tight recommendation memo, behind the climate samples. A blueprint sits next to the implementation notes that stay off the menu, so it does not lead the site.",
        space_after=1,
    )

    heading(doc, "Leave off the site")
    bullet(
        doc,
        "Avon supplier diversity and inclusion strategy and roll-out plan. ",
        "A client strategy from about 2019–2020, written at WEConnect International, with named Avon and WEConnect staff and supplier findings. Both organisations would have to clear it in writing. Even then it is the wrong lead sample for this practice.",
    )
    bullet(
        doc,
        "RentRemote deal blueprint. ",
        "Marked confidential. It names a real company, a workstation count, a distribution plan, and partners. It stays private unless AstroFinance and RentRemote both agree in writing. Their deal does not become a carbon case study.",
    )
    bullet(
        doc,
        "Aircraft APU deal blueprint. ",
        "Marked “Confidential · For Discussion Purposes Only.” Meridian Aero Parts is labeled fictional. Honeywell, ETHZilla, STORM Partners, and the leasing structure are real-world references, and those names are the point of the piece. It fails the same test as SimAxis, with less left once the names come out.",
    )
    bullet(
        doc,
        "A New Chessboard (2012). ",
        "Copyright Changing Our World, Inc. The about-the-authors note lists Nicolas Picard as the 2012 Corporate Social Investment Fellow. It does not list him as the author of a chapter. A one-line fellowship credit on LinkedIn is the right use. The sustainability chapter is fourteen years old.",
    )
    bullet(
        doc,
        "Strategic CSR at the Business Core. ",
        "The brief speaks as IG Advisors and closes with a Liquidnet case. The examples are early-2010s. A release from IG Advisors would still leave the wrong lead sample on the site.",
        space_after=1,
    )

    heading(doc, "Rewrite from scratch, into supplier-carbon software")
    para(
        doc,
        "These are new pieces. They copy a shape. They do not revise the old PDF.",
        space_after=2,
    )
    bullet(
        doc,
        "Avon’s method, as a new briefing. ",
        "That plan assesses a program across policy, planning, processes and standards, measurement, tracking and reporting, and accountability, then sets a 6–12 month roll-out. Write a new practitioner briefing on how a supplier-carbon program rolls out. No Avon or WEConnect names, no spend, no supplier findings. Advertise it as a practitioner briefing.",
    )
    bullet(
        doc,
        "The CSR brief’s shape, as a new briefing. ",
        "That brief moves through buyer behaviour, brand, cost (including local sourcing), and multi-stakeholder collaboration, then a short case. Write a new practitioner briefing on what a supplier-carbon product has to show a buyer. No Liquidnet, no IG Advisors, no 2010 figures. Advertise it as a practitioner briefing.",
        space_after=1,
    )

    heading(doc, "Order on the site, once the new pieces exist")
    bullet(doc, "First. ", "The AI briefing and the public paper on where supplier-carbon programs break.")
    bullet(
        doc,
        "Second, after a written yes from AstroFinance. ",
        "The three short pieces, or a cut of white-paper sections 6.1 to 6.3. Label it a product white paper.",
    )
    bullet(doc, "Optional, below the fold. ", "A link to the published AmCham article, dated about 2014.")
    bullet(
        doc,
        "Optional, and last. ",
        "The cut SimAxis memo, labeled a client memo, after the same written yes.",
        space_after=2,
    )
    para(
        doc,
        "If AstroFinance, the AmCham publisher, IG Advisors, Avon, or WEConnect has already said yes in writing, name the file. The line for that file is the only one that moves.",
        size=9,
        italic=True,
        color=MUTED,
        space_after=0,
    )
    doc.save("/workspace/Nicolas-Picard-Writing-Samples-To-Share.docx")


if __name__ == "__main__":
    build()
