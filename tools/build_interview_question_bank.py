#!/usr/bin/env python3
"""Merge the two interview prep documents into one structured question bank.

Every line of question/answer text is copied verbatim (run by run, so inline
bold/italic survives) from the source .docx files by paragraph index. The only
text this script introduces is the front matter, the section/subsection titles,
and the "[No answer drafted yet]" placeholders.

Usage:  python3 tools/build_interview_question_bank.py
Outputs: Interview-Question-Bank.docx and Interview-Question-Bank.md
"""

from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

REPO = Path(__file__).resolve().parent.parent
SRC_A = REPO / "sources" / "Interview-questions.docx"
SRC_B = REPO / "sources" / "solutions-consultant-prep.docx"
OUT_DOCX = REPO / "Interview-Question-Bank.docx"
OUT_MD = REPO / "Interview-Question-Bank.md"

BULLET_GLYPHS = [("\u2022", "Calibri"), ("\u25e6", "Calibri"), ("\u25aa", "Calibri")]
PLACEHOLDER = "[No answer drafted yet]"


def load(path):
    return docx.Document(str(path)).paragraphs


def ilvl_of(par):
    pr = par._p.pPr
    if pr is None or pr.numPr is None or pr.numPr.ilvl is None:
        return 0
    return pr.numPr.ilvl.val


# --------------------------------------------------------------------------
# Document plan
#
# Ops:
#   ("h1"|"h2"|"h3"|"h4", literal text)
#   ("hsrc", level, src, idx)        heading taken verbatim from a source line
#   ("q", src, idx)                  question line rendered as bold paragraph
#   ("p", src, idx)                  prose paragraph
#   ("b", src, idx, level)           single bullet at an explicit level
#   ("rng", src, first, last, shift) bullet range, nesting taken from source
#   ("ph",)                          placeholder for an unanswered question
#   ("intro", literal text)          front-matter paragraph
#   ("toc", [items])                 contents list
# --------------------------------------------------------------------------

PLAN = [
    ("title", "Interview Question Bank"),
    ("intro", "A single, consolidated prep document merging \u201cInterview questions\u201d and "
              "\u201csolutions consultant prep\u201d. All question and answer text is reproduced "
              "verbatim from those two documents; only the section structure is new."),
    ("toc", [
        "1. General interview questions",
        "2. Solutions consulting & solutions engineering questions",
        "3. Product manager questions",
        "4. Technical & engineering collaboration questions",
        "5. B2B SaaS questions",
        "6. Sustainability, ESG & value chain questions",
        "7. Experience bank: examples and evidence",
        "Appendix: Questions still to answer",
    ]),

    # ---------------------------------------------------------------- 1
    ("h1", "1. General interview questions"),

    ("h2", "1.1 Background and experience"),
    ("hsrc", 3, "A", 1),
    ("ph",),
    ("hsrc", 3, "A", 4),
    ("ph",),
    ("hsrc", 3, "A", 6),
    ("rng", "A", 7, 14, 0),
    ("hsrc", 3, "A", 16),
    ("rng", "A", 17, 24, 0),
    ("hsrc", 3, "A", 26),
    ("rng", "A", 27, 31, 0),

    ("h2", "1.2 Teamwork"),
    ("hsrc", 3, "A", 91),
    ("rng", "A", 93, 99, 0),

    ("h2", "1.3 Project & stakeholder management"),
    ("hsrc", 3, "A", 84),
    ("rng", "A", 85, 88, 0),
    ("hsrc", 3, "A", 312),
    ("rng", "A", 313, 316, 0),

    # ---------------------------------------------------------------- 2
    ("h1", "2. Solutions consulting & solutions engineering questions"),

    ("h2", "2.1 Discovery & building client relationships"),
    ("hsrc", 3, "B", 129),
    ("rng", "B", 130, 136, 1),
    ("hsrc", 3, "B", 149),
    ("rng", "B", 150, 159, 1),

    ("h2", "2.2 Demos & explaining complex concepts"),
    ("hsrc", 3, "A", 363),
    ("rng", "A", 364, 365, 0),
    ("hsrc", 3, "B", 88),
    ("p", "B", 89),
    ("p", "B", 90),
    ("hsrc", 3, "A", 374),
    ("ph",),
    ("hsrc", 3, "A", 372),
    ("ph",),

    ("h2", "2.3 Technical scoping, implementation & troubleshooting"),
    ("hsrc", 3, "B", 91),
    ("p", "B", 92),
    ("p", "B", 93),
    ("p", "B", 94),
    ("hsrc", 3, "B", 95),
    ("p", "B", 96),
    ("p", "B", 97),
    ("p", "B", 98),
    ("p", "B", 99),
    ("rng", "B", 100, 102, 2),
    ("p", "B", 103),
    ("hsrc", 3, "B", 104),
    ("rng", "B", 105, 110, 1),

    ("h2", "2.4 Objections, pushback & saying no"),
    ("hsrc", 3, "B", 111),
    ("rng", "B", 112, 120, 1),
    ("hsrc", 3, "A", 370),
    ("ph",),
    ("hsrc", 3, "A", 367),
    ("rng", "A", 368, 368, 0),
    ("hsrc", 3, "A", 376),
    ("ph",),
    ("hsrc", 3, "A", 378),
    ("rng", "A", 379, 381, 0),

    ("h2", "2.5 Managing client expectations & complex requirements"),
    ("hsrc", 3, "B", 122),
    ("rng", "B", 123, 128, 1),
    ("hsrc", 3, "B", 138),
    ("b", "B", 139, 0),
    ("b", "B", 140, 1),
    ("b", "B", 141, 0),
    ("b", "B", 142, 1),
    ("b", "B", 143, 0),
    ("b", "B", 144, 1),
    ("b", "B", 145, 0),
    ("b", "B", 146, 1),
    ("b", "B", 147, 0),
    ("b", "B", 148, 1),

    ("h2", "2.6 Cross-functional collaboration"),
    ("hsrc", 3, "B", 160),
    ("rng", "B", 161, 166, 1),

    ("h2", "2.7 Supporting the sales cycle & competitive positioning"),
    ("hsrc", 3, "B", 197),
    ("rng", "B", 198, 219, 0),
    ("hsrc", 3, "B", 220),
    ("rng", "B", 221, 223, 1),
    ("hsrc", 3, "B", 224),
    ("rng", "B", 225, 229, 1),

    ("h2", "2.8 Customer success, retention & scaling"),
    ("hsrc", 3, "B", 167),
    ("hsrc", 4, "B", 168),
    ("rng", "B", 169, 170, 0),
    ("hsrc", 4, "B", 171),
    ("rng", "B", 172, 173, 0),
    ("hsrc", 4, "B", 174),
    ("rng", "B", 175, 177, 0),
    ("hsrc", 3, "B", 179),
    ("hsrc", 4, "B", 180),
    ("rng", "B", 181, 182, 0),
    ("hsrc", 4, "B", 183),
    ("rng", "B", 184, 185, 0),
    ("hsrc", 4, "B", 186),
    ("rng", "B", 187, 188, 0),
    ("hsrc", 4, "B", 189),
    ("rng", "B", 190, 193, 0),
    ("hsrc", 4, "B", 194),
    ("rng", "B", 195, 196, 0),

    # ---------------------------------------------------------------- 3
    ("h1", "3. Product manager questions"),

    ("h2", "3.1 Motivation & ways of working"),
    ("hsrc", 3, "A", 33),
    ("rng", "A", 34, 36, 0),
    ("hsrc", 3, "A", 38),
    ("rng", "A", 39, 40, 0),

    ("h2", "3.2 Strategy"),
    ("hsrc", 3, "A", 42),
    ("rng", "A", 43, 55, 0),
    ("p", "A", 57),
    ("p", "A", 58),
    ("hsrc", 3, "A", 60),
    ("rng", "A", 61, 66, 0),
    ("p", "A", 68),
    ("p", "A", 69),
    ("p", "A", 71),
    ("p", "A", 73),
    ("hsrc", 3, "A", 75),
    ("rng", "A", 76, 80, 0),

    ("h2", "3.3 Design"),
    ("hsrc", 3, "A", 102),
    ("rng", "A", 104, 107, 0),
    ("hsrc", 3, "A", 109),
    ("rng", "A", 110, 116, 0),
    ("hsrc", 3, "A", 118),
    ("rng", "A", 119, 135, 0),
    ("hsrc", 3, "A", 137),
    ("rng", "A", 138, 142, 0),

    ("h2", "3.4 Develop & deliver"),
    ("hsrc", 3, "A", 145),
    ("rng", "A", 146, 149, 0),
    ("hsrc", 3, "A", 151),
    ("rng", "A", 152, 156, 0),
    ("hsrc", 3, "A", 158),
    ("rng", "A", 159, 166, 0),
    ("hsrc", 3, "A", 168),
    ("rng", "A", 169, 170, 0),

    ("h2", "3.5 Product performance & analytics"),
    ("hsrc", 3, "A", 173),
    ("rng", "A", 174, 179, 0),
    ("hsrc", 3, "A", 181),
    ("rng", "A", 182, 187, 0),
    ("hsrc", 3, "A", 189),
    ("rng", "A", 190, 200, 0),

    ("h2", "3.6 How-would-you product questions"),
    ("hsrc", 3, "A", 207),
    ("p", "A", 208),
    ("rng", "A", 209, 214, 0),
    ("hsrc", 3, "A", 216),
    ("ph",),
    ("hsrc", 3, "A", 219),
    ("rng", "A", 220, 224, 0),
    ("hsrc", 3, "A", 226),
    ("rng", "A", 227, 230, 0),

    # ---------------------------------------------------------------- 4
    ("h1", "4. Technical & engineering collaboration questions"),
    ("hsrc", 3, "A", 282),
    ("rng", "A", 283, 288, 0),
    ("hsrc", 3, "A", 290),
    ("rng", "A", 291, 293, 0),
    ("hsrc", 3, "A", 295),
    ("rng", "A", 296, 301, 0),
    ("hsrc", 3, "A", 303),
    ("rng", "A", 304, 310, 0),
    ("hsrc", 3, "A", 318),
    ("rng", "A", 319, 326, 0),
    ("hsrc", 3, "A", 328),
    ("rng", "A", 329, 339, 0),
    ("hsrc", 3, "A", 341),
    ("rng", "A", 342, 353, 0),
    ("hsrc", 3, "A", 355),
    ("rng", "A", 356, 359, 0),

    # ---------------------------------------------------------------- 5
    ("h1", "5. B2B SaaS questions"),
    ("hsrc", 3, "A", 235),
    ("rng", "A", 236, 246, 0),
    ("hsrc", 3, "A", 248),
    ("rng", "A", 249, 260, 0),
    ("hsrc", 3, "A", 262),
    ("rng", "A", 263, 273, 0),

    # ---------------------------------------------------------------- 6
    ("h1", "6. Sustainability, ESG & value chain questions"),
    ("hsrc", 3, "A", 384),
    ("rng", "A", 385, 394, 0),
    ("hsrc", 3, "A", 396),
    ("q", "A", 397),
    ("rng", "A", 398, 404, 0),
    ("hsrc", 3, "A", 406),
    ("q", "A", 407),
    ("rng", "A", 408, 408, 0),
    ("hsrc", 3, "A", 410),
    ("q", "A", 411),
    ("rng", "A", 412, 418, 0),
    ("hsrc", 3, "A", 420),
    ("q", "A", 421),
    ("rng", "A", 422, 422, 0),
    ("hsrc", 3, "A", 424),
    ("q", "A", 425),
    ("rng", "A", 426, 426, 0),
    ("hsrc", 3, "A", 428),
    ("q", "A", 429),
    ("rng", "A", 430, 437, 0),
    ("hsrc", 3, "A", 439),
    ("q", "A", 440),
    ("rng", "A", 441, 445, 0),
    ("hsrc", 3, "A", 447),
    ("q", "A", 448),
    ("rng", "A", 449, 454, 0),
    ("hsrc", 3, "A", 456),
    ("q", "A", 457),
    ("rng", "A", 458, 463, 0),

    # ---------------------------------------------------------------- 7
    ("h1", "7. Experience bank: examples and evidence"),
    ("hsrc", 2, "B", 0),
    ("rng", "B", 1, 27, 1),
    ("hsrc", 2, "B", 29),
    ("rng", "B", 30, 46, 0),
    ("hsrc", 2, "B", 48),
    ("rng", "B", 49, 84, 0),

    # ---------------------------------------------------------------- appendix
    ("h1", "Appendix: Questions still to answer"),
    ("b", "A", 1, 0),
    ("b", "A", 4, 0),
    ("b", "A", 216, 0),
    ("b", "A", 370, 0),
    ("b", "A", 372, 0),
    ("b", "A", 374, 0),
    ("b", "A", 376, 0),
]


def build():
    src = {"A": load(SRC_A), "B": load(SRC_B)}
    doc = docx.Document()
    style_document(doc)
    num_id = define_bullet_list(doc)
    md = []

    def copy_runs(target, par, force_plain=False):
        for run in par.runs:
            new = target.add_run(run.text)
            if not force_plain:
                new.bold = run.bold
                new.italic = run.italic
                new.underline = run.underline

    def md_text(par, force_plain=False):
        out = []
        for run in par.runs:
            text = run.text
            if not force_plain and run.bold and text.strip():
                lead = text[: len(text) - len(text.lstrip())]
                trail = text[len(text.rstrip()):]
                text = f"{lead}**{text.strip()}**{trail}"
            out.append(text)
        return "".join(out) or par.text

    def heading(level, text):
        doc.add_heading(text, level=level)
        md.append(("#" * level + " " + text, True))

    def bullet(par, level):
        level = max(0, min(level, len(BULLET_GLYPHS) - 1))
        target = doc.add_paragraph(style="List Paragraph")
        apply_bullet(target, num_id, level)
        copy_runs(target, par)
        md.append(("  " * level + "- " + md_text(par), False))

    for op in PLAN:
        kind = op[0]

        if kind == "title":
            doc.add_heading(op[1], level=0)
            md.append(("# " + op[1], True))
        elif kind == "intro":
            par = doc.add_paragraph(op[1])
            par.runs[0].italic = True
            md.append(("_" + op[1] + "_", True))
        elif kind == "toc":
            doc.add_heading("Contents", level=1)
            md.append(("## Contents", True))
            for item in op[1]:
                par = doc.add_paragraph(item, style="List Paragraph")
                apply_bullet(par, num_id, 0)
                md.append(("- " + item, False))
        elif kind in ("h1", "h2", "h3", "h4"):
            heading(int(kind[1]), op[1])
        elif kind == "hsrc":
            _, level, s, idx = op
            # Strip stray markdown emphasis markers that came through in a few
            # of the source headings; the wording itself is untouched.
            text = src[s][idx].text.strip().replace("**", "")
            doc.add_heading(text, level=level)
            md.append(("#" * level + " " + text, True))
        elif kind == "q":
            _, s, idx = op
            par = src[s][idx]
            target = doc.add_paragraph()
            copy_runs(target, par)
            for run in target.runs:
                run.bold = True
            md.append(("**" + par.text.strip() + "**", True))
        elif kind == "p":
            _, s, idx = op
            par = src[s][idx]
            target = doc.add_paragraph()
            copy_runs(target, par)
            md.append((md_text(par), True))
        elif kind == "b":
            _, s, idx, level = op
            bullet(src[s][idx], level)
        elif kind == "rng":
            _, s, first, last, shift = op
            for idx in range(first, last + 1):
                par = src[s][idx]
                if not par.text.strip():
                    continue
                bullet(par, ilvl_of(par) - shift)
        elif kind == "ph":
            target = doc.add_paragraph()
            run = target.add_run(PLACEHOLDER)
            run.italic = True
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            md.append(("_" + PLACEHOLDER + "_", True))
        else:
            raise ValueError(f"unknown op {kind}")

    add_page_numbers(doc)
    doc.save(str(OUT_DOCX))

    lines = []
    for text, blank_after in md:
        if text.startswith("#") and lines and lines[-1] != "":
            lines.append("")
        lines.append(text)
        if blank_after:
            lines.append("")
    OUT_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")


def style_document(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before in (
        ("Heading 1", 20, RGBColor(0x1F, 0x36, 0x4D), 24),
        ("Heading 2", 15, RGBColor(0x2E, 0x5B, 0x7A), 16),
        ("Heading 3", 12, RGBColor(0x1F, 0x1F, 0x1F), 14),
        ("Heading 4", 11, RGBColor(0x40, 0x40, 0x40), 10),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.font.italic = False
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    listpar = doc.styles["List Paragraph"]
    listpar.paragraph_format.space_after = Pt(3)
    listpar.paragraph_format.line_spacing = 1.15


def define_bullet_list(doc):
    """Add a three level bullet definition and return its numId."""
    numbering = doc.part.numbering_part.element
    abstract_id = 900
    num_id = 900

    abstract = docx.oxml.OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    for level, (glyph, font) in enumerate(BULLET_GLYPHS):
        lvl = docx.oxml.OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        for tag, attr, value in (
            ("w:start", "w:val", "1"),
            ("w:numFmt", "w:val", "bullet"),
            ("w:lvlText", "w:val", glyph),
            ("w:lvlJc", "w:val", "left"),
        ):
            node = docx.oxml.OxmlElement(tag)
            node.set(qn(attr), value)
            lvl.append(node)

        ppr = docx.oxml.OxmlElement("w:pPr")
        ind = docx.oxml.OxmlElement("w:ind")
        ind.set(qn("w:left"), str(360 + 360 * level))
        ind.set(qn("w:hanging"), "270")
        ppr.append(ind)
        lvl.append(ppr)

        rpr = docx.oxml.OxmlElement("w:rPr")
        fonts = docx.oxml.OxmlElement("w:rFonts")
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(attr), font)
        fonts.set(qn("w:hint"), "default")
        rpr.append(fonts)
        lvl.append(rpr)

        abstract.append(lvl)
    numbering.insert(0, abstract)

    num = docx.oxml.OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = docx.oxml.OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def apply_bullet(par, num_id, level):
    ppr = par._p.get_or_add_pPr()
    numpr = docx.oxml.OxmlElement("w:numPr")
    ilvl = docx.oxml.OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num = docx.oxml.OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    numpr.append(ilvl)
    numpr.append(num)
    ppr.append(numpr)
    par.paragraph_format.left_indent = Pt((360 + 360 * level) / 20)
    par.paragraph_format.first_line_indent = Pt(-270 / 20)


def add_page_numbers(doc):
    footer = doc.sections[0].footer
    par = footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    for kind, text in (("begin", None), ("instr", "PAGE"), ("end", None)):
        if kind == "instr":
            node = docx.oxml.OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = text
        else:
            node = docx.oxml.OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        run._r.append(node)


if __name__ == "__main__":
    build()
    print(f"wrote {OUT_DOCX.name} and {OUT_MD.name}")
